"""資料の形式ごとの検証と差分（design.md §3.9）。

MD・CSV・PPTXの原本を読み、形式の誤り（errors）・注意（warnings）・中身の概要（info）と、前の版との差分（diff）を返す。
検証の条件（CSVの必須列・キー列・型等）は管理側が持つ値（`checks`）だけを使い、候補ファイルの中身からは変えさせない。
ファイルの中の外部リンク先へは通信しない。結果はJSONで保存し、画面とエージェントにそのまま見せる。
"""

import csv
import datetime
import difflib
import hashlib
import io
import itertools
import posixpath
import re
import zipfile
from collections import Counter
from collections.abc import Iterable
from typing import Any

from django.utils.translation import gettext as _
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

MD = "md"
CSV = "csv"
PPTX = "pptx"
OTHER = "other"
FORMATS = (MD, CSV, PPTX, OTHER)
EXTENSIONS = {".md": MD, ".markdown": MD, ".csv": CSV, ".pptx": PPTX}
# 差分の行・スライドの変わり方
DIFF_CHANGES = ("added", "removed", "changed")

CSV_ENCODINGS = ("utf-8", "utf-8-sig", "cp932")
CSV_DELIMITERS = {",": "comma", "\t": "tab", ";": "semicolon"}
CSV_DELIMITER_CHOICES = tuple(CSV_DELIMITERS)
COLUMN_TYPES = ("string", "integer", "decimal", "date")

# 差分に載せる量の上限。超えた分は数だけを数え、truncated にする
MAX_DIFF_LINES = 400
MAX_DIFF_ROWS = 50
MAX_DIFF_SLIDES = 50
MAX_TEXT = 300
MAX_EXAMPLES = 5
MAX_COLUMNS = 100
# MDの差分を取る大きさの上限（difflibは長い文書で遅くなる）
MAX_DIFF_BYTES = 2 * 1024 * 1024
# PPTX（zip）の展開の上限。小さいファイルが大きく展開されるもの（zip bomb）を読まない
MAX_ZIP_ENTRIES = 5000
MAX_ZIP_UNCOMPRESSED = 300 * 1024 * 1024
# 編集した資料で変わってはいけない要素（python-pptxで正しく編集できないもの）
UNEDITABLE_ELEMENTS = ("picture", "chart", "graphic", "media")

INTEGER_RE = re.compile(r"-?\d+")
DECIMAL_RE = re.compile(r"-?\d+(\.\d+)?")
DATE_RE = re.compile(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})")
LEADING_ZERO_RE = re.compile(r"0\d+")


def format_of(name: str) -> str:
    return EXTENSIONS.get(posixpath.splitext(name.lower())[1], OTHER)


def normalize_checks(checks: dict[str, Any] | None) -> dict[str, Any]:
    """検証の条件の既定値を補う。条件はCSVだけが持つ（encoding・delimiterの空文字は自動で判別する）。
    列の型は {"column", "type"} の並び（列名を辞書のキーにしない。APIでキーの書き方を変換するため）"""
    checks = checks or {}
    types = {t["column"]: t["type"] for t in checks.get("column_types") or [] if t.get("type", "string") != "string"}
    return {
        "encoding": checks.get("encoding") or "",
        "delimiter": checks.get("delimiter") or "",
        "required_columns": list(checks.get("required_columns") or []),
        "key_columns": list(checks.get("key_columns") or []),
        "column_types": [{"column": column, "type": kind} for column, kind in types.items()],
    }


def _column_types(checks: dict[str, Any]) -> dict[str, str]:
    return {t["column"]: t["type"] for t in checks["column_types"]}


def inspect(
    data: bytes,
    name: str,
    *,
    checks: dict[str, Any] | None = None,
    before: tuple[bytes, str] | None = None,
    edited: bool = False,
) -> dict[str, Any]:
    """資料を検証し、前の版（before：bytesとファイル名）との差分を取る。

    edited は、エージェントが元の資料を編集して作った候補か（添付・アップロードの原本のままの置換ではないか）。
    編集した候補では、正しく編集できない要素が変わっていないことも確かめる（PPTX）
    """
    result = _Result(format_of(name), edited)
    previous = before if before is not None and format_of(before[1]) == result.format else None
    if before is not None and previous is None:
        result.warn(_("形式が前の版（%(name)s）と違います") % {"name": before[1]})
    checks = normalize_checks(checks)
    if result.format == MD:
        _inspect_md(result, data, previous[0] if previous else None)
    elif result.format == CSV:
        _inspect_csv(result, data, checks, previous[0] if previous else None)
    elif result.format == PPTX:
        _inspect_pptx(result, data, previous[0] if previous else None)
    elif before is not None:
        result.diff = {"identical": data == before[0]}
    if result.diff is not None and previous is not None:
        result.diff["identical"] = data == previous[0]
    return result.as_dict()


def diff_counts(validation: dict[str, Any]) -> dict[str, int]:
    """差分の数だけ（SkillChangeの記録と、エージェントへの短い報告に使う）"""
    diff = validation.get("diff") or {}
    return {k: v for k, v in diff.items() if isinstance(v, int) and not isinstance(v, bool)}


def describe(validation: dict[str, Any]) -> str:
    """検証の結果と差分の要約（エージェント向けの1〜数行）"""
    if not validation:
        return "not checked"
    parts = []
    if counts := diff_counts(validation):
        parts.append("changes: " + ", ".join(f"{k}={v}" for k, v in counts.items()))
    elif (validation.get("diff") or {}).get("identical"):
        parts.append("no change from the published version")
    errors = validation.get("errors") or []
    warnings = validation.get("warnings") or []
    parts.append(f"errors: {'; '.join(errors)}" if errors else "no errors")
    if warnings:
        parts.append(f"warnings: {'; '.join(warnings)}")
    return " / ".join(parts)


class _Result:
    def __init__(self, fmt: str, edited: bool):
        self.format = fmt
        self.edited = edited
        self.errors: list[str] = []
        self.warnings: list[str] = []
        self.info: dict[str, Any] = {}
        self.diff: dict[str, Any] | None = None

    def error(self, message: str) -> None:
        self.errors.append(message)

    def warn(self, message: str) -> None:
        self.warnings.append(message)

    def as_dict(self) -> dict[str, Any]:
        return {
            "format": self.format,
            "edited": self.edited,
            "errors": self.errors,
            "warnings": self.warnings,
            "info": self.info,
            "diff": self.diff,
        }


def _clip(value: str | None) -> str | None:
    if value is None or len(value) <= MAX_TEXT:
        return value
    return value[: MAX_TEXT - 1] + "…"


def _examples(items: Iterable[str]) -> str:
    items = list(items)
    text = "、".join(items[:MAX_EXAMPLES])
    return text + (_(" ほか%(n)s件") % {"n": len(items) - MAX_EXAMPLES} if len(items) > MAX_EXAMPLES else "")


# --- テキスト ---


def _decode(data: bytes, encoding: str = "") -> tuple[str, str] | None:
    """文字列と、使った文字コード。指定がなければ UTF-8（BOM付きを含む）、Shift_JIS（cp932）の順に試す"""
    if encoding:
        candidates = [encoding]
    elif data.startswith(b"\xef\xbb\xbf"):
        candidates = ["utf-8-sig"]
    else:
        candidates = ["utf-8", "cp932"]
    for candidate in candidates:
        try:
            return data.decode(candidate), candidate
        except UnicodeDecodeError:
            continue
    return None


def _newline(text: str) -> str:
    crlf = text.count("\r\n")
    lf = text.count("\n") - crlf
    if crlf and lf:
        return "mixed"
    return "crlf" if crlf else "lf"


def _inspect_md(result: _Result, data: bytes, before: bytes | None) -> None:
    if b"\x00" in data:
        result.error(_("テキストのファイルではありません"))
        return
    decoded = _decode(data)
    if decoded is None:
        result.error(_("文字コードを判別できません（UTF-8かShift_JISで保存してください）"))
        return
    text, encoding = decoded
    result.info = {"encoding": encoding, "newline": _newline(text), "lines": len(text.splitlines())}
    if before is None:
        return
    previous = _decode(before)
    if previous is None:
        return
    before_text, before_encoding = previous
    if before_encoding != encoding:
        result.warn(
            _("文字コードが変わりました（%(before)s → %(after)s）") % {"before": before_encoding, "after": encoding}
        )
    if _newline(before_text) != result.info["newline"]:
        result.warn(_("改行の形式が変わりました"))
    if before == data:
        result.diff = {"identical": True, "lines_added": 0, "lines_removed": 0}
        return
    if len(before) > MAX_DIFF_BYTES or len(data) > MAX_DIFF_BYTES:
        result.warn(_("ファイルが大きいため、本文の差分は表示しません"))
        return
    lines = list(
        difflib.unified_diff(before_text.splitlines(), text.splitlines(), "before", "after", n=2, lineterm="")
    )[2:]
    added = sum(1 for line in lines if line.startswith("+"))
    removed = sum(1 for line in lines if line.startswith("-"))
    result.diff = {
        "lines_added": added,
        "lines_removed": removed,
        "text": [_clip(line) or "" for line in lines[:MAX_DIFF_LINES]],
        "truncated": len(lines) > MAX_DIFF_LINES,
    }


# --- CSV ---


class _Table:
    def __init__(self, encoding: str, delimiter: str, header: list[str], rows: list[tuple[int, list[str]]]):
        self.encoding = encoding
        self.delimiter = delimiter
        self.header = header
        self.rows = rows  # (行番号, 値)。値はすべて文字列のまま扱い、数値に変えない（先頭の0を落とさない）

    def records(self) -> list[tuple[int, dict[str, str]]]:
        return [(line, dict(zip(self.header, values, strict=False))) for line, values in self.rows]


def _sniff_delimiter(text: str) -> str:
    """見出しの行で、引用符の外にいちばん多く現れる区切り文字（引用符の中の改行も見出しの続きとして扱う）"""
    counts = dict.fromkeys(CSV_DELIMITERS, 0)
    quoted = False
    for char in text:
        if char == '"':
            quoted = not quoted  # "" は閉じてすぐ開くので、引用の中のままになる
        elif not quoted and char == "\n":
            break
        elif not quoted and char in counts:
            counts[char] += 1
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] else ","


def _read_csv(data: bytes, checks: dict[str, Any], result: _Result | None) -> _Table | None:
    """CSVを読む。読めなければ誤りを result に足してNone（前の版を読むときは result を渡さない）"""

    def fail(message: str) -> None:
        if result is not None:
            result.error(message)

    decoded = _decode(data, checks["encoding"])
    if decoded is None:
        fail(
            _("文字コード（%(encoding)s）として読めません") % {"encoding": checks["encoding"]}
            if checks["encoding"]
            else _("文字コードを判別できません（UTF-8かShift_JISで保存してください）")
        )
        return None
    text, encoding = decoded
    delimiter = checks["delimiter"] or _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    rows: list[tuple[int, list[str]]] = []
    try:
        for values in reader:
            if values:  # 空の行は読み飛ばす
                rows.append((reader.line_num, values))
    except csv.Error as exc:
        fail(_("CSVとして読めません（%(line)s行目：%(error)s）") % {"line": reader.line_num, "error": exc})
        return None
    if not rows:
        fail(_("CSVに見出しの行がありません"))
        return None
    header = [h.strip() for h in rows[0][1]]
    return _Table(encoding, delimiter, header, rows[1:])


def _type_ok(kind: str, value: str) -> bool:
    if kind == "integer":
        return bool(INTEGER_RE.fullmatch(value))
    if kind == "decimal":
        return bool(DECIMAL_RE.fullmatch(value))
    if kind == "date":
        match = DATE_RE.fullmatch(value)
        if not match:
            return False
        try:
            datetime.date(*(int(g) for g in match.groups()))
        except ValueError:
            return False
    return True


def _check_table(result: _Result, table: _Table, checks: dict[str, Any]) -> None:
    header = table.header
    if any(not h for h in header):
        result.error(_("見出しが空の列があります"))
    if duplicated := [h for h, n in Counter(header).items() if n > 1 and h]:
        result.error(_("同じ見出しの列があります：%(columns)s") % {"columns": _examples(duplicated)})
    if bad := [str(line) for line, values in table.rows if len(values) != len(header)]:
        result.error(_("列の数が見出しと違う行があります：%(lines)s行目") % {"lines": _examples(bad)})
    if missing := [c for c in checks["required_columns"] if c not in header]:
        result.error(_("必須の列がありません：%(columns)s") % {"columns": _examples(missing)})
    keys = checks["key_columns"]
    records = table.records()
    if missing_keys := [c for c in keys if c not in header]:
        result.error(_("キーの列がありません：%(columns)s") % {"columns": _examples(missing_keys)})
    elif keys:
        empty = [str(line) for line, record in records if any(not record.get(c, "") for c in keys)]
        if empty:
            result.error(_("キーが空の行があります：%(lines)s行目") % {"lines": _examples(empty)})
        counted = Counter(_key(record, keys) for _line, record in records)
        if duplicated_keys := [k for k, n in counted.items() if n > 1]:
            result.error(
                _("キーが重複しています：%(keys)s") % {"keys": _examples(_key_text(k) for k in duplicated_keys)}
            )
    for column, kind in _column_types(checks).items():
        if column not in header or kind == "string":
            continue
        wrong = [
            _("%(line)s行目「%(value)s」") % {"line": line, "value": _clip(record[column])}
            for line, record in records
            if record.get(column) and not _type_ok(kind, record[column])
        ]
        if wrong:
            result.error(
                _("列「%(column)s」に%(kind)sでない値があります：%(values)s")
                % {"column": column, "kind": _type_label(kind), "values": _examples(wrong)}
            )


def _type_label(kind: str) -> str:
    return {"integer": _("整数"), "decimal": _("数値"), "date": _("日付")}.get(kind, kind)


def _key(record: dict[str, str], keys: list[str]) -> tuple[str, ...]:
    """キーの列の値の組。比べるときは組のまま使い、表示するときだけ _key_text() で文字列にする"""
    return tuple(record.get(c, "") for c in keys)


def _key_text(key: tuple[str, ...]) -> str:
    return " / ".join(key)


def _inspect_csv(result: _Result, data: bytes, checks: dict[str, Any], before: bytes | None) -> None:
    table = _read_csv(data, checks, result)
    if table is None:
        return
    result.info = {
        "encoding": table.encoding,
        "delimiter": CSV_DELIMITERS.get(table.delimiter, table.delimiter),
        "columns": table.header[:MAX_COLUMNS],
        "rows": len(table.rows),
    }
    _check_table(result, table, checks)
    if before is None:
        return
    # 前の版は文字コード・区切りを判別して読む（条件を変えた後でも、前の版を前の形のまま読めるように）
    previous = _read_csv(before, {**checks, "encoding": "", "delimiter": ""}, None)
    if previous is None:
        return
    if previous.encoding != table.encoding:
        result.warn(
            _("文字コードが変わりました（%(before)s → %(after)s）")
            % {"before": previous.encoding, "after": table.encoding}
        )
    if previous.delimiter != table.delimiter:
        result.warn(_("区切り文字が変わりました"))
    _check_leading_zeros(result, previous, table, checks)
    result.diff = _diff_tables(previous, table, checks["key_columns"])


def _check_leading_zeros(result: _Result, before: _Table, after: _Table, checks: dict[str, Any]) -> None:
    """前の版に先頭が0の値（コード等）があった列から、先頭の0が消えていないか（表計算ソフトでの保存で起きやすい）"""
    numeric = {c for c, kind in _column_types(checks).items() if kind in ("integer", "decimal")}
    before_records = [r for _l, r in before.records()]
    after_records = [r for _l, r in after.records()]
    for column in after.header:
        if column in numeric or column not in before.header:
            continue
        zeros = {r[column] for r in before_records if LEADING_ZERO_RE.fullmatch(r.get(column, ""))}
        if not zeros:
            continue
        # 0付きの値が消え、0を取った値が新しく現れたもの（行を消しただけ・別の行の値と取り違えない）だけを数える
        values = {r.get(column, "") for r in after_records}
        previous = {r.get(column, "") for r in before_records}
        examples = sorted(
            v
            for v in zeros
            if v not in values and (stripped := v.lstrip("0") or "0") in values and stripped not in previous
        )
        if examples:
            result.error(
                _(
                    "列「%(column)s」の先頭の0が消えています（前の版の例：「%(example)s」）。値を文字列のまま保存してください"
                )
                % {"column": column, "example": examples[0]}
            )


def _diff_tables(before: _Table, after: _Table, keys: list[str]) -> dict[str, Any]:
    columns = [c for c in after.header if c in before.header]
    diff: dict[str, Any] = {
        "columns_added": [c for c in after.header if c not in before.header][:MAX_COLUMNS],
        "columns_removed": [c for c in before.header if c not in after.header][:MAX_COLUMNS],
    }
    # 件数はすべて数え、差分の行は上限の数だけ組み立てる（大きいCSVでメモリを使いすぎない）
    rows: list[dict[str, Any]] = []
    if keys and all(k in before.header and k in after.header for k in keys):
        old = {_key(r, keys): (line, r) for line, r in before.records()}
        new = {_key(r, keys): (line, r) for line, r in after.records()}
        # 足した・消した列の値も、同じキーの行の変更として数える（ない列は空として比べる）
        compared = [*after.header, *(c for c in before.header if c not in after.header)]
        added = [k for k in new if k not in old]
        removed = [k for k in old if k not in new]
        changed = sorted(
            (new[k][0], k)
            for k in new.keys() & old.keys()
            if any(old[k][1].get(c, "") != new[k][1].get(c, "") for c in compared)
        )
        for k in added[:MAX_DIFF_ROWS]:
            line, record = new[k]
            rows.append(
                {"change": "added", "key": _key_text(k), "line": line, "cells": _cells(record, after.header, "after")}
            )
        for k in removed[: MAX_DIFF_ROWS - len(rows)]:
            line, record = old[k]
            rows.append(
                {
                    "change": "removed",
                    "key": _key_text(k),
                    "line": line,
                    "cells": _cells(record, before.header, "before"),
                }
            )
        for line, k in changed[: MAX_DIFF_ROWS - len(rows)]:
            a, b = old[k][1], new[k][1]
            cells = [
                {"column": c, "before": _clip(a.get(c)), "after": _clip(b.get(c))}
                for c in compared
                if a.get(c, "") != b.get(c, "")
            ]
            rows.append({"change": "changed", "key": _key_text(k), "line": line, "cells": cells})
        counts = {"rows_added": len(added), "rows_removed": len(removed), "rows_changed": len(changed)}
    else:
        # キーがなければ、行の組合せの増減だけを数える
        old_rows = Counter(tuple(r.get(c, "") for c in columns) for _l, r in before.records())
        new_rows = Counter(tuple(r.get(c, "") for c in columns) for _l, r in after.records())
        added_rows, removed_rows = new_rows - old_rows, old_rows - new_rows
        for change, counter in (("added", added_rows), ("removed", removed_rows)):
            side = "after" if change == "added" else "before"
            for values in itertools.islice(counter.elements(), MAX_DIFF_ROWS - len(rows)):
                cells = _cells(dict(zip(columns, values, strict=True)), columns, side)
                rows.append({"change": change, "key": "", "line": None, "cells": cells})
        counts = {"rows_added": added_rows.total(), "rows_removed": removed_rows.total(), "rows_changed": 0}
    total = sum(counts.values())
    diff |= counts
    diff |= {"rows": rows, "truncated": total > len(rows)}
    return diff


def _cells(record: dict[str, str], columns: list[str], side: str) -> list[dict[str, str | None]]:
    other = "before" if side == "after" else "after"
    return [{"column": c, side: _clip(record.get(c, "")), other: None} for c in columns[:MAX_COLUMNS]]


# --- PPTX ---


SHAPE_KINDS = {MSO_SHAPE_TYPE.PICTURE: "picture", MSO_SHAPE_TYPE.MEDIA: "media", MSO_SHAPE_TYPE.GROUP: "group"}


def _shape_kind(shape) -> str:
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    try:
        kind = SHAPE_KINDS.get(shape.shape_type)
    except NotImplementedError:
        kind = None
    if kind:
        return kind
    # SmartArt・OLE等は、グラフ・表でないGraphicFrame
    return "graphic" if shape.__class__.__name__ == "GraphicFrame" else "shape"


def _shape_text(shape) -> str | None:
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return shape.text_frame.text
    if getattr(shape, "has_table", False) and shape.has_table:
        return "\n".join("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
    return None


def _walk(shapes, prefix: str = ""):
    for shape in shapes:
        yield prefix, shape
        if _shape_kind(shape) == "group":
            yield from _walk(shape.shapes, f"{prefix}{shape.shape_id}/")


def _fingerprint(shape, kind: str) -> str:
    """正しく編集できない要素の中身のhash。数が同じでも、差し替えたことが分かるように比べる"""
    try:
        if kind == "picture":
            return shape.image.sha1
        if kind == "chart":
            return hashlib.sha256(shape.chart_part.blob).hexdigest()
    except (AttributeError, KeyError, ValueError):
        pass
    # SmartArt・OLE・動画等は、要素のXML（参照する部品のIDを含む）で比べる
    return hashlib.sha256(shape._element.xml.encode()).hexdigest()


def _slide_structure(slide) -> dict[str, Any]:
    texts: dict[str, dict[str, str]] = {}
    elements: Counter[str] = Counter()
    fingerprints: list[str] = []
    for prefix, shape in _walk(slide.shapes):
        kind = _shape_kind(shape)
        elements[kind] += 1
        if kind in UNEDITABLE_ELEMENTS:
            fingerprints.append(f"{kind}:{_fingerprint(shape, kind)}")
        text = _shape_text(shape)
        if text is not None:
            texts[f"{prefix}{shape.shape_id}"] = {"name": shape.name, "text": text}
    placeholders = sorted(f"{p.placeholder_format.idx}:{p.placeholder_format.type}" for p in slide.placeholders)
    return {
        "id": slide.slide_id,
        "layout": slide.slide_layout.name,
        "placeholders": placeholders,
        "texts": texts,
        "elements": dict(elements),
        "fingerprints": sorted(fingerprints),
    }


def _read_pptx(data: bytes, result: _Result | None) -> dict[str, Any] | None:
    def fail(message: str) -> None:
        if result is not None:
            result.error(message)

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile:
        fail(_("PPTXとして読めません（壊れているか、PPTXではありません）"))
        return None
    if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_ZIP_UNCOMPRESSED:
        fail(_("展開すると大きすぎるため、読みません"))
        return None
    try:
        presentation = Presentation(io.BytesIO(data))
        slides = [_slide_structure(slide) for slide in presentation.slides]
        layouts = [layout.name for master in presentation.slide_masters for layout in master.slide_layouts]
    except Exception as exc:  # python-pptxは壊れたファイルで様々な例外を出す
        fail(_("PPTXとして読めません（%(error)s）") % {"error": exc.__class__.__name__})
        return None
    return {"slides": slides, "layouts": layouts}


def _inspect_pptx(result: _Result, data: bytes, before: bytes | None) -> None:
    deck = _read_pptx(data, result)
    if deck is None:
        return
    result.info = {
        "slides": len(deck["slides"]),
        "layouts": deck["layouts"][:MAX_COLUMNS],
        "elements": dict(sum((Counter(s["elements"]) for s in deck["slides"]), Counter())),
    }
    if before is None:
        return
    previous = _read_pptx(before, None)
    if previous is None:
        return
    result.diff = _diff_decks(previous, deck)
    if result.edited:
        _check_edited_deck(result, previous, deck)


def _diff_decks(before: dict[str, Any], after: dict[str, Any]) -> dict[str, Any]:
    old = {s["id"]: (i, s) for i, s in enumerate(before["slides"], 1)}
    new = {s["id"]: (i, s) for i, s in enumerate(after["slides"], 1)}
    entries: list[dict[str, Any]] = []
    changed = 0
    for slide_id, (index, slide) in new.items():
        if slide_id not in old:
            texts = [{"shape": t["name"], "before": None, "after": _clip(t["text"])} for t in slide["texts"].values()]
            entries.append(
                {
                    "slide": index,
                    "change": "added",
                    "layout_before": None,
                    "layout_after": slide["layout"],
                    "texts": texts,
                    "elements": [],
                }
            )
            continue
        _old_index, previous = old[slide_id]
        texts = []
        for key in previous["texts"].keys() | slide["texts"].keys():
            a = previous["texts"].get(key)
            b = slide["texts"].get(key)
            if (a and a["text"]) != (b and b["text"]):
                texts.append(
                    {
                        "shape": (b or a or {}).get("name", key),
                        "before": _clip(a and a["text"]),
                        "after": _clip(b and b["text"]),
                    }
                )
        elements = [
            {"kind": kind, "before": previous["elements"].get(kind, 0), "after": slide["elements"].get(kind, 0)}
            for kind in sorted(previous["elements"].keys() | slide["elements"].keys())
            if previous["elements"].get(kind, 0) != slide["elements"].get(kind, 0)
        ]
        layout_changed = previous["layout"] != slide["layout"] or previous["placeholders"] != slide["placeholders"]
        if texts or elements or layout_changed:
            changed += 1
            entries.append(
                {
                    "slide": index,
                    "change": "changed",
                    "layout_before": previous["layout"],
                    "layout_after": slide["layout"],
                    "texts": sorted(texts, key=lambda t: t["shape"])[:MAX_DIFF_ROWS],
                    "elements": elements,
                }
            )
    for slide_id, (index, slide) in old.items():
        if slide_id not in new:
            entries.append(
                {
                    "slide": index,
                    "change": "removed",
                    "layout_before": slide["layout"],
                    "layout_after": None,
                    "texts": [],
                    "elements": [],
                }
            )
    order = [s for s in new if s in old]
    return {
        "slides_added": sum(1 for s in new if s not in old),
        "slides_removed": sum(1 for s in old if s not in new),
        "slides_changed": changed,
        "reordered": order != [s for s in old if s in new],
        "slides": entries[:MAX_DIFF_SLIDES],
        "truncated": len(entries) > MAX_DIFF_SLIDES,
    }


def _check_edited_deck(result: _Result, before: dict[str, Any], after: dict[str, Any]) -> None:
    """編集した候補で、正しく編集できない要素（画像・グラフ・SmartArt・動画等）とレイアウトが変わっていないか。
    変わっていれば編集した資料として公開させず、原本のままの置換を案内する（design.md §3.9）"""
    if Counter(before["layouts"]) != Counter(after["layouts"]):
        result.error(_("スライドのレイアウトが変わっています。レイアウトの変更は編集では扱えません"))
    old = {s["id"]: s for s in before["slides"]}
    broken = []
    for index, slide in enumerate(after["slides"], 1):
        previous = old.get(slide["id"])
        if previous is None:
            continue
        if previous["fingerprints"] != slide["fingerprints"]:
            broken.append(str(index))
    if broken:
        result.error(
            _(
                "%(slides)s枚目のスライドで、画像・グラフ・SmartArt等が変わっています。これらは編集では扱えないため、"
                "完成したファイルをアップロードして置き換えてください"
            )
            % {"slides": _examples(broken)}
        )
