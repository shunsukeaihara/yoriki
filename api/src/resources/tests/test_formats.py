"""資料の形式ごとの検証と差分（design.md §3.9）"""

import io
import zipfile

from pptx import Presentation
from pptx.util import Inches

from resources import formats

KEYED = {"key_columns": ["code"]}


def _csv(text: str, encoding: str = "utf-8") -> bytes:
    return text.encode(encoding)


def _deck(*texts: str, picture: bool = False) -> bytes:
    presentation = Presentation()
    for text in texts:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = text
    if picture:
        slide = presentation.slides[0]
        png = (
            b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4"
            b"\x89\x00\x00\x00\rIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
        )
        slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1))
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def _edit_title(data: bytes, index: int, text: str, *, drop_pictures: bool = False) -> bytes:
    presentation = Presentation(io.BytesIO(data))
    slide = presentation.slides[index]
    slide.shapes.title.text = text
    if drop_pictures:
        for shape in list(slide.shapes):
            if shape.shape_type == 13:  # PICTURE
                shape._element.getparent().remove(shape._element)
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


class TestMarkdown:
    def test_diff_and_encoding(self):
        before = "# ルール\n- 支店は区別しない\n".encode()
        after = "# ルール\n- 支店ごとに区別する\n".encode("cp932")
        result = formats.inspect(after, "rules.md", before=(before, "rules.md"))
        assert result["format"] == "md"
        assert result["errors"] == []
        assert result["info"]["encoding"] == "cp932"
        assert any("文字コードが変わりました" in w for w in result["warnings"])
        assert result["diff"]["lines_added"] == 1
        assert result["diff"]["lines_removed"] == 1
        assert "+- 支店ごとに区別する" in result["diff"]["text"]

    def test_binary_is_an_error(self):
        result = formats.inspect(b"\x00\x01", "rules.md")
        assert result["errors"]


class TestCsv:
    def test_checks(self):
        data = _csv("code,name,amount,date\n001,A,10,2026-01-31\n001,B,x,2026-02-30\n,C,1,\n002,D\n")
        checks = {
            "required_columns": ["name", "kana"],
            "key_columns": ["code"],
            "column_types": [{"column": "amount", "type": "integer"}, {"column": "date", "type": "date"}],
        }
        result = formats.inspect(data, "vendors.csv", checks=checks)
        errors = "\n".join(result["errors"])
        assert "列の数が見出しと違う行があります：5行目" in errors
        assert "必須の列がありません：kana" in errors
        assert "キーが空の行があります：4行目" in errors
        assert "キーが重複しています：001" in errors
        assert "列「amount」に整数でない値があります：3行目「x」" in errors
        assert "列「date」に日付でない値があります：3行目「2026-02-30」" in errors
        assert result["info"]["rows"] == 4
        assert result["info"]["columns"] == ["code", "name", "amount", "date"]

    def test_diff_by_key(self):
        before = _csv("code,name\n001,A\n002,B\n003,C\n")
        after = _csv("code,name,kana\n001,A,エー\n002,Bee,ビー\n004,D,ディー\n")
        result = formats.inspect(after, "vendors.csv", checks=KEYED, before=(before, "vendors.csv"))
        diff = result["diff"]
        assert result["errors"] == []
        assert (diff["rows_added"], diff["rows_removed"], diff["rows_changed"]) == (1, 1, 1)
        assert diff["columns_added"] == ["kana"]
        changed = next(r for r in diff["rows"] if r["change"] == "changed")
        assert changed["key"] == "002"
        assert changed["cells"] == [{"column": "name", "before": "B", "after": "Bee"}]

    def test_composite_keys_and_large_diffs(self, monkeypatch):
        # 値に「 / 」を含んでも、別のキーとして扱う
        data = _csv('a,b\n"x / y",z\nx,"y / z"\n')
        assert formats.inspect(data, "v.csv", checks={"key_columns": ["a", "b"]})["errors"] == []
        monkeypatch.setattr(formats, "MAX_DIFF_ROWS", 2)
        before = _csv("code,name\n" + "".join(f"{i:03},A\n" for i in range(10)))
        after = _csv("code,name\n" + "".join(f"{i:03},B\n" for i in range(10)))
        diff = formats.inspect(after, "v.csv", checks=KEYED, before=(before, "v.csv"))["diff"]
        assert diff["rows_changed"] == 10
        assert len(diff["rows"]) == 2
        assert diff["truncated"] is True

    def test_diff_without_key_counts_rows(self):
        before = _csv("code,name\n001,A\n002,B\n")
        after = _csv("code,name\n001,A\n002,B\n002,B\n")
        diff = formats.inspect(after, "v.csv", before=(before, "v.csv"))["diff"]
        assert (diff["rows_added"], diff["rows_removed"]) == (1, 0)

    def test_lost_leading_zeros_are_an_error(self):
        before = _csv("code,name\n001,A\n010,B\n")
        after = _csv("code,name\n1,A\n10,B\n")
        result = formats.inspect(after, "v.csv", before=(before, "v.csv"))
        assert any("先頭の0が消えています" in e for e in result["errors"])

    def test_removing_rows_is_not_a_lost_zero(self):
        before = _csv("code,name\n001,A\n2,B\n")
        after = _csv("code,name\n2,B\n")
        result = formats.inspect(after, "v.csv", before=(before, "v.csv"))
        assert result["errors"] == []
        assert result["diff"]["identical"] is False
        # 0を取った値が別の行に元からあれば、行を消しただけとして扱う
        before = _csv("code,name\n001,A\n1,B\n")
        assert formats.inspect(_csv("code,name\n1,B\n"), "v.csv", before=(before, "v.csv"))["errors"] == []

    def test_previous_version_is_read_by_its_own_encoding(self):
        before = _csv("code,name\n001,東京\n", "cp932")
        after = _csv("code,name\n001,東京本社\n")
        checks = {"encoding": "utf-8", "key_columns": ["code"]}
        result = formats.inspect(after, "v.csv", checks=checks, before=(before, "v.csv"))
        assert result["errors"] == []
        assert any("文字コードが変わりました" in w for w in result["warnings"])
        assert result["diff"]["rows_changed"] == 1

    def test_shift_jis_and_explicit_encoding(self):
        data = _csv("コード,名前\n001,東京\n", "cp932")
        assert formats.inspect(data, "v.csv")["info"]["encoding"] == "cp932"
        result = formats.inspect(data, "v.csv", checks={"encoding": "utf-8"})
        assert result["errors"] == ["文字コード（utf-8）として読めません"]
        # 引用符の中のカンマは区切りとして数えない
        quoted = formats.inspect(_csv('"姓, 名";code\n"山田, 太郎";001\n'), "v.csv")
        assert quoted["info"]["delimiter"] == "semicolon"
        assert quoted["info"]["columns"] == ["姓, 名", "code"]
        tab = formats.inspect(_csv("a\tb\n1\t2\n"), "v.csv")
        assert tab["info"]["delimiter"] == "tab"
        assert tab["info"]["columns"] == ["a", "b"]


class TestPptx:
    def test_structure_and_text_diff(self):
        before = _deck("表紙", "概要")
        after = _edit_title(before, 1, "概要（改訂）")
        result = formats.inspect(after, "deck.pptx", before=(before, "deck.pptx"), edited=True)
        assert result["errors"] == []
        assert result["info"]["slides"] == 2
        diff = result["diff"]
        assert (diff["slides_added"], diff["slides_removed"], diff["slides_changed"]) == (0, 0, 1)
        [slide] = diff["slides"]
        assert slide["slide"] == 2
        assert slide["texts"][0]["before"] == "概要"
        assert slide["texts"][0]["after"] == "概要（改訂）"

    def test_edited_deck_must_keep_pictures(self):
        before = _deck("表紙", picture=True)
        after = _edit_title(before, 0, "新しい表紙", drop_pictures=True)
        edited = formats.inspect(after, "deck.pptx", before=(before, "deck.pptx"), edited=True)
        assert any("画像・グラフ・SmartArt" in e for e in edited["errors"])
        # 原本のままの置換（アップロード）なら、要素が変わっても誤りにしない
        replaced = formats.inspect(after, "deck.pptx", before=(before, "deck.pptx"))
        assert replaced["errors"] == []
        assert replaced["diff"]["slides"][0]["elements"] == [{"kind": "picture", "before": 1, "after": 0}]

    def test_edited_deck_must_keep_the_same_pictures(self):
        before = _deck("表紙", picture=True)
        presentation = Presentation(io.BytesIO(before))
        slide = presentation.slides[0]
        [picture] = [shape for shape in slide.shapes if shape.shape_type == 13]
        picture._element.getparent().remove(picture._element)
        other = (
            b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x02\x00\x00\x00\x02\x08\x02\x00\x00\x00\xfd\xd4\x9as"
            b"\x00\x00\x00\x0fIDATx\x9cc\xf8\xcf\xc0\x00\x02\x0c\x0c\x00\x00\x0e\x00\x01\xe6\xa0\x9fV\x00\x00\x00\x00IEND"
            b"\xaeB`\x82"
        )
        slide.shapes.add_picture(io.BytesIO(other), Inches(1), Inches(1))
        out = io.BytesIO()
        presentation.save(out)
        # 画像の数が同じでも、差し替えれば誤り
        result = formats.inspect(out.getvalue(), "deck.pptx", before=(before, "deck.pptx"), edited=True)
        assert any("画像・グラフ・SmartArt" in e for e in result["errors"])

    def test_broken_and_huge_files(self, monkeypatch):
        assert formats.inspect(b"not a zip", "deck.pptx")["errors"]
        monkeypatch.setattr(formats, "MAX_ZIP_UNCOMPRESSED", 1024)
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("big.bin", b"\0" * 1025)
        assert formats.inspect(out.getvalue(), "deck.pptx")["errors"] == ["展開すると大きすぎるため、読みません"]


def test_other_formats_are_only_compared():
    result = formats.inspect(b"%PDF", "a.pdf", before=(b"%PDF", "a.pdf"))
    assert result["format"] == "other"
    assert result["diff"] == {"identical": True}
    assert formats.inspect(b"a,b\n", "a.csv", before=(b"#", "a.md"))["warnings"]
