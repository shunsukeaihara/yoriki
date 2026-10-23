"""形式に合わせて読む（harness.md §7.8 `read`）。大きいファイルは、まず全体の構造を返し、範囲を指定して読ませる"""

import base64
import csv
import io
import subprocess
from pathlib import Path
from typing import Any

from .common import ToolError, check_size, resolve, sha256_of
from .text import MAX_LINES, decode, is_binary, page_of

# PDFの既定で読むページ数と、1回で返す画像の数
DEFAULT_PDF_PAGES = 10
MAX_IMAGES = 5
# モデルへ渡す画像の長辺（これより大きい画像は縮小する）
MAX_IMAGE_SIDE = 1568
# PNGがこれより大きければJPEGにする（写真・ノイズの多いスキャン。各社の1枚の上限（Anthropicは5MB）に収める）
MAX_PNG_BYTES = 1024 * 1024
JPEG_QUALITY = 85
# Excelを1回で表に出すセルの数と、指定できる範囲の大きさ（読む前に確かめる）
MAX_CELLS = 4000
MAX_RANGE_CELLS = 50_000

SUFFIX_KIND = {
    ".pdf": "pdf",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".docx": "docx",
    ".pptx": "pptx",
    ".csv": "csv",
    ".tsv": "csv",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".webp": "image",
    ".bmp": "image",
}


def _int(request: dict[str, Any], name: str, default: int, minimum: int = 1) -> int:
    value = request.get(name, default)
    if value is None:
        return default
    if not isinstance(value, int) or isinstance(value, bool) or value < minimum:
        raise ToolError("invalid_argument", f"{name} must be an integer >= {minimum}.")
    return value


def run_read(request: dict[str, Any]) -> dict[str, Any]:
    path = resolve(request.get("path"))
    if path.is_dir():
        raise ToolError("not_a_file", f"{path} is a directory. Use glob to list it.")
    if not path.is_file():
        raise ToolError("not_a_file", f"{path} is not a regular file.")
    size = check_size(path)
    kind = SUFFIX_KIND.get(path.suffix.lower(), "text")
    base = {"ok": True, "path": str(path), "size": size, "sha256": sha256_of(path)}
    images = bool(request.get("images", False))
    readers = {
        "pdf": lambda: read_pdf(path, request.get("pages"), images),
        "xlsx": lambda: read_xlsx(path, request.get("sheet"), request.get("range")),
        "docx": lambda: _paged(read_docx(path), request, "docx"),
        "pptx": lambda: _paged(read_pptx(path), request, "pptx"),
        "image": lambda: read_image(path, images),
    }
    if kind in readers:
        return base | readers[kind]()
    data = path.read_bytes()
    if is_binary(data):
        return base | {"kind": "binary", "mime": _mime(path), "text": _binary_notice(path)}
    decoded = decode(data)
    offset, limit = _int(request, "offset", 1), min(_int(request, "limit", MAX_LINES), MAX_LINES)
    body, meta = page_of(decoded.text, offset, limit)
    result = base | meta | {"encoding": decoded.encoding, "bom": decoded.bom, "newline": _newline(decoded.newline)}
    if kind == "csv":
        result |= {"kind": "csv", "csv": _csv_summary(decoded.text, path)}
    else:
        result["kind"] = "text"
    result["text"] = body if body else "(empty file)"
    return result


def _newline(value: str) -> str:
    return {"\r\n": "crlf", "\n": "lf"}.get(value, "mixed")


def _paged(text: str, request: dict[str, Any], kind: str) -> dict[str, Any]:
    offset, limit = _int(request, "offset", 1), min(_int(request, "limit", MAX_LINES), MAX_LINES)
    body, meta = page_of(text, offset, limit)
    return {"kind": kind, "text": body or "(no text)"} | meta


def _mime(path: Path) -> str:
    try:
        return subprocess.run(
            ["file", "--brief", "--mime-type", str(path)], capture_output=True, text=True, check=True, timeout=10
        ).stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return "application/octet-stream"


def _binary_notice(path: Path) -> str:
    return f"{path} is a binary file ({_mime(path)}). Use shell with python to inspect it."


def _csv_summary(text: str, path: Path) -> dict[str, Any]:
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    columns = next(reader, [])
    rows = sum(1 for _ in reader)
    return {"delimiter": delimiter, "columns": columns, "rows": rows}


# --- PDF ---


def _parse_pages(spec: Any, total: int) -> list[int]:
    if spec in (None, ""):
        return list(range(1, min(total, DEFAULT_PDF_PAGES) + 1))
    if not isinstance(spec, str):
        raise ToolError("invalid_argument", 'pages must be a string such as "1-5" or "3,7".')
    pages: list[int] = []
    for raw in spec.split(","):
        part = raw.strip()
        try:
            if "-" in part:
                first, last = (int(x) for x in part.split("-", 1))
            else:
                first = last = int(part)
        except ValueError as exc:
            raise ToolError("invalid_argument", f'Invalid pages "{spec}". Use a form such as "1-5" or "3,7".') from exc
        if first < 1 or last < first:
            raise ToolError("invalid_argument", f'Invalid page range "{part}".')
        pages.extend(range(first, min(last, total) + 1))
    return sorted(set(pages))


def read_pdf(path: Path, spec: Any, images: bool) -> dict[str, Any]:
    import pdfplumber  # noqa: PLC0415  使うときだけ読む（起動を軽くする）

    try:
        pdf = pdfplumber.open(path)
    except Exception as exc:
        raise ToolError("unreadable", f"The PDF could not be opened: {exc}") from exc
    with pdf:
        total = len(pdf.pages)
        pages = _parse_pages(spec, total)
        out, scanned = [], []
        for number in pages:
            text = (pdf.pages[number - 1].extract_text() or "").strip()
            if not text:
                scanned.append(number)
                text = "(no text layer on this page; it must be read as an image)"
            out.append(f"--- page {number} / {total} ---\n{text}")
    result: dict[str, Any] = {"kind": "pdf", "total_pages": total, "pages": pages, "text": "\n\n".join(out)}
    if scanned:
        result["pages_without_text"] = scanned
    if images and scanned:
        result["images"] = _render_pdf_pages(path, scanned[:MAX_IMAGES])
    if pages and pages[-1] < total:
        result["next_pages"] = f"{pages[-1] + 1}-{min(pages[-1] + DEFAULT_PDF_PAGES, total)}"
    return result


def _render_pdf_pages(path: Path, pages: list[int]) -> list[dict[str, Any]]:
    import pypdfium2  # noqa: PLC0415

    document = pypdfium2.PdfDocument(path)
    rendered = []
    try:
        for number in pages:
            page = document[number - 1]
            width, height = page.get_size()
            scale = min(MAX_IMAGE_SIDE / max(width, height), 3.0)
            image = page.render(scale=scale).to_pil()
            rendered.append({"page": number} | _encode_image(image))
    finally:
        document.close()
    return rendered


# --- 画像 ---


def _encode_image(image: Any) -> dict[str, Any]:
    image.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    if image.mode not in ("RGB", "L"):
        image = image.convert("RGB")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    mime = "image/png"
    if buffer.tell() > MAX_PNG_BYTES:
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=JPEG_QUALITY)
        mime = "image/jpeg"
    return {
        "mime": mime,
        "width": image.width,
        "height": image.height,
        "data": base64.b64encode(buffer.getvalue()).decode(),
    }


def read_image(path: Path, images: bool) -> dict[str, Any]:
    from PIL import Image, UnidentifiedImageError  # noqa: PLC0415

    try:
        with Image.open(path) as image:
            result: dict[str, Any] = {
                "kind": "image",
                "format": image.format,
                "width": image.width,
                "height": image.height,
                "text": f"Image {image.format} {image.width}x{image.height}",
            }
            if images:
                image.load()
                result["images"] = [_encode_image(image.copy())]
    except (UnidentifiedImageError, OSError) as exc:
        raise ToolError("unreadable", f"The image could not be read: {exc}") from exc
    return result


# --- Excel ---


def _bounds(ws: Any, target: str) -> tuple[int, int, int, int]:
    """範囲の端を求め、セルを読む前に大きさを確かめる（巨大な範囲でメモリを使い切らないように）"""
    from openpyxl.utils import range_boundaries  # noqa: PLC0415

    try:
        min_col, min_row, max_col, max_row = range_boundaries(target)
    except (ValueError, TypeError) as exc:
        raise ToolError("invalid_argument", f'Invalid range "{target}". Use a form such as A1:F50.') from exc
    min_col, min_row = min_col or 1, min_row or 1
    max_col, max_row = max_col or ws.max_column, max_row or ws.max_row
    if (max_col - min_col + 1) * (max_row - min_row + 1) > MAX_RANGE_CELLS:
        raise ToolError(
            "too_large",
            f'Range "{target}" has more than {MAX_RANGE_CELLS} cells. Pass a smaller range, such as A1:F200.',
        )
    return min_col, min_row, max_col, max_row


def read_xlsx(path: Path, sheet: Any, cell_range: Any) -> dict[str, Any]:
    from openpyxl import load_workbook  # noqa: PLC0415

    try:
        formulas = load_workbook(path, data_only=False)
        values = load_workbook(path, data_only=True)
    except Exception as exc:
        raise ToolError("unreadable", f"The workbook could not be opened: {exc}") from exc
    sheets = [
        {"name": ws.title, "dimensions": ws.dimensions, "rows": ws.max_row, "columns": ws.max_column} for ws in formulas
    ]
    if sheet is None:
        lines = [f"{s['name']}: {s['dimensions']} ({s['rows']} rows x {s['columns']} columns)" for s in sheets]
        text = "Sheets:\n" + "\n".join(lines) + "\n\nPass sheet (and range, such as A1:F50) to read cells."
        return {"kind": "xlsx", "sheets": sheets, "text": text}
    if not isinstance(sheet, str) or sheet not in formulas.sheetnames:
        raise ToolError("not_found", f'Sheet "{sheet}" does not exist. Sheets: {", ".join(formulas.sheetnames)}')
    ws_f, ws_v = formulas[sheet], values[sheet]
    target = cell_range if isinstance(cell_range, str) and cell_range else ws_f.dimensions
    min_col, min_row, max_col, max_row = _bounds(ws_f, target)
    out, cells = [], 0
    rows = ws_f.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
    for row in rows:
        parts = []
        for cell in row:
            if cell.value is None:
                continue
            if cells >= MAX_CELLS:
                break
            cells += 1
            value = ws_v[cell.coordinate].value
            if isinstance(cell.value, str) and cell.value.startswith("="):
                parts.append(f"{cell.coordinate}={value!r} ⟨{cell.value}⟩")
            else:
                parts.append(f"{cell.coordinate}={cell.value!r}")
        if parts:
            out.append(" | ".join(parts))
        if cells >= MAX_CELLS:
            out.append(f"[stopped after {MAX_CELLS} cells; pass a smaller range]")
            break
    text = f"Sheet {sheet}, range {target}\n" + ("\n".join(out) if out else "(no values in this range)")
    return {"kind": "xlsx", "sheets": sheets, "sheet": sheet, "range": target, "text": text}


# --- Word ---


def _md_table(rows: list[list[str]]) -> list[str]:
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    clean = [[c.replace("\n", " ").replace("|", "\\|") for c in r] for r in rows]
    lines = ["| " + " | ".join(clean[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in clean[1:]]
    return lines


def read_docx(path: Path) -> str:
    from docx import Document  # noqa: PLC0415
    from docx.table import Table  # noqa: PLC0415
    from docx.text.paragraph import Paragraph  # noqa: PLC0415

    try:
        document = Document(str(path))
    except Exception as exc:
        raise ToolError("unreadable", f"The document could not be opened: {exc}") from exc
    lines: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name if block.style is not None else "") or ""
            if style.startswith("Heading") and style[-1:].isdigit():
                lines.append("#" * int(style[-1]) + " " + text)
            else:
                lines.append(text)
        elif isinstance(block, Table):
            lines += _md_table([[cell.text for cell in row.cells] for row in block.rows])
    return "\n".join(lines)


# --- PowerPoint ---


def read_pptx(path: Path) -> str:
    from pptx import Presentation  # noqa: PLC0415

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise ToolError("unreadable", f"The presentation could not be opened: {exc}") from exc
    lines: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines.append(f"## Slide {number} (layout: {slide.slide_layout.name})")
        for shape in slide.shapes:
            label = shape.name
            if shape.is_placeholder:
                label += f" [placeholder: {shape.placeholder_format.type}]"
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(f"- {label}: {shape.text_frame.text.strip()}")
            elif shape.has_table:
                lines.append(f"- {label}: table")
                lines += _md_table([[cell.text for cell in row.cells] for row in shape.table.rows])
            elif shape.is_placeholder:
                lines.append(f"- {label}: (empty)")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append(f"Notes: {slide.notes_slide.notes_text_frame.text.strip()}")
    return "\n".join(lines)
