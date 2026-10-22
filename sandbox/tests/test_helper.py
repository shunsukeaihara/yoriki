"""補助コマンド `yoriki-sandbox` のサブコマンドと入出力（CONTRACT.md §4）"""

import json
import subprocess
import time

import pytest

# Office・PDFの見本は、Sandboxの中のライブラリで作る（ホストにライブラリを入れないため）
MAKE_SAMPLES = r"""
python - <<'PY'
from openpyxl import Workbook
from docx import Document
from pptx import Presentation
from PIL import Image
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

wb = Workbook(); ws = wb.active; ws.title = "明細"
ws.append(["品名", "数量", "単価", "金額"])
ws.append(["りんご", 3, 100, "=B2*C2"])
ws.append(["みかん", 5, 80, "=B3*C3"])
wb.create_sheet("空")
wb.save("/workspace/work/samples/見積.xlsx")

doc = Document(); doc.add_heading("契約書", 1); doc.add_paragraph("第1条 目的")
t = doc.add_table(rows=2, cols=2); t.cell(0, 0).text = "甲"; t.cell(0, 1).text = "乙"
t.cell(1, 0).text = "株式会社A"; t.cell(1, 1).text = "株式会社B"
doc.save("/workspace/work/samples/契約.docx")

prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "売上報告"; s.placeholders[1].text = "前年比 120%"
s.notes_slide.notes_text_frame.text = "補足を話す"
prs.save("/workspace/work/samples/報告.pptx")

fig = plt.figure(); fig.text(0.1, 0.5, "Invoice total 12345"); fig.savefig("/workspace/work/samples/text.pdf")
Image.new("RGB", (2000, 1000), "white").save("/workspace/work/samples/scan.pdf")
Image.new("RGB", (3000, 1500), "red").save("/workspace/work/samples/写真.png")
PY
"""


@pytest.fixture(scope="module")
def samples(sandbox):
    sandbox.sh("mkdir -p /workspace/work/samples && " + MAKE_SAMPLES)
    return "/workspace/work/samples"


def test_version(sandbox):
    version = sandbox.helper("version", {})
    assert (version["ok"], version["contract"]) == (True, 1)
    assert version["python"].startswith("3.13.")


class TestExec:
    def test_reports_the_process_group_and_exit_code(self, sandbox):
        result = sandbox.exec("echo out; echo err >&2; exit 7")
        assert result.exit_code == 7
        # 標準出力と標準エラーは、どちらもコマンドの出力として流れる
        assert result.output == "out\nerr\n"
        assert result.events[0]["event"] == "started"
        assert result.events[0]["pgid"] > 1
        assert result.events[-1] == {"event": "exited", "exit_code": 7, "signal": None}

    def test_stdin_is_closed_and_state_does_not_carry_over(self, sandbox):
        assert sandbox.exec("read x; echo [$?]").output == "[1]\n"
        sandbox.sh("cd /tmp; export FOO=1")
        assert sandbox.sh("pwd; echo [${FOO:-}]") == "/workspace/work\n[]\n"

    def test_environment_is_only_what_is_passed(self, sandbox):
        env = dict(line.split("=", 1) for line in sandbox.sh("env").splitlines() if "=" in line)
        assert env["HOME"] == "/workspace/work/.home"
        assert env["PATH"].startswith("/workspace/env/bin:")
        assert "YORIKI_SECRET_KEY" not in env
        assert sandbox.sh("command -v python").strip() == "/workspace/env/bin/python"

    def test_kill_stops_the_whole_group(self, sandbox):
        started = time.monotonic()
        # サブプロセスごと止まることを確かめる（グループへsignalを送る）
        proc = subprocess.Popen(
            ["docker", "exec", "-i", sandbox.name, "/opt/yoriki/bin/yoriki-sandbox", "exec"],
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        assert proc.stdin and proc.stderr
        proc.stdin.write(json.dumps({"command": "sleep 300 & sleep 300; wait", "cwd": "/workspace/work", "env": {}}))
        proc.stdin.close()
        started_event = json.loads(proc.stderr.readline())
        killed = sandbox.helper("kill", {"pgid": started_event["pgid"], "signal": "TERM"})
        assert killed == {"ok": True, "killed": True}
        proc.wait(timeout=30)
        assert time.monotonic() - started < 30
        assert sandbox.exec('pgrep -f "[s]leep 300"').exit_code == 1

    def test_rejects_a_bad_request(self, sandbox):
        result = sandbox.exec("true", cwd="/etc")
        assert result.exit_code == 125
        assert '"outside_workspace"' in result.output


class TestRead:
    def test_text_with_line_numbers_and_paging(self, sandbox):
        result = sandbox.helper("read", {"path": "/workspace/input/メモ.txt", "offset": 2, "limit": 1})
        assert result["kind"] == "text"
        assert result["text"] == "     2│二行目"
        assert result["total_lines"] == 3
        assert result["next_offset"] == 3
        assert len(result["sha256"]) == 64

    def test_shift_jis_csv(self, sandbox):
        result = sandbox.helper("read", {"path": "/workspace/input/請求書一覧.csv"})
        assert result["kind"] == "csv"
        assert result["encoding"] == "cp932"
        assert result["newline"] == "crlf"
        assert result["csv"] == {"delimiter": ",", "columns": ["番号", "取引先", "金額"], "rows": 2}
        assert "髙橋商店" in result["text"]

    def test_binary(self, sandbox):
        result = sandbox.helper("read", {"path": "/workspace/input/data.bin"})
        assert result["kind"] == "binary"
        assert "python" in result["text"]

    def test_pdf_pages_and_scanned_pages(self, sandbox, samples):
        result = sandbox.helper("read", {"path": f"{samples}/text.pdf"})
        assert result["total_pages"] == 1
        assert "Invoice total 12345" in result["text"]
        scanned = sandbox.helper("read", {"path": f"{samples}/scan.pdf", "images": True})
        assert scanned["pages_without_text"] == [1]
        assert scanned["images"][0]["mime"] == "image/png"
        assert max(scanned["images"][0]["width"], scanned["images"][0]["height"]) <= 1568

    def test_xlsx_sheets_then_cells_with_formulas(self, sandbox, samples):
        overview = sandbox.helper("read", {"path": f"{samples}/見積.xlsx"})
        assert [s["name"] for s in overview["sheets"]] == ["明細", "空"]
        cells = sandbox.helper("read", {"path": f"{samples}/見積.xlsx", "sheet": "明細", "range": "A1:D3"})
        assert "A2='りんご'" in cells["text"]
        assert "⟨=B2*C2⟩" in cells["text"]
        missing = sandbox.helper("read", {"path": f"{samples}/見積.xlsx", "sheet": "なし"})
        assert missing["error"]["code"] == "not_found"

    def test_docx_and_pptx(self, sandbox, samples):
        docx = sandbox.helper("read", {"path": f"{samples}/契約.docx"})
        assert "# 契約書" in docx["text"]
        assert "| 株式会社A | 株式会社B |" in docx["text"]
        pptx = sandbox.helper("read", {"path": f"{samples}/報告.pptx"})
        assert "Slide 1" in pptx["text"]
        assert "売上報告" in pptx["text"]
        assert "Notes: 補足を話す" in pptx["text"]

    def test_image_is_resized(self, sandbox, samples):
        meta = sandbox.helper("read", {"path": f"{samples}/写真.png"})
        assert (meta["width"], meta["height"]) == (3000, 1500)
        assert "images" not in meta
        image = sandbox.helper("read", {"path": f"{samples}/写真.png", "images": True})["images"][0]
        assert (image["width"], image["height"]) == (1568, 784)

    def test_errors(self, sandbox):
        assert sandbox.helper("read", {"path": "/workspace/work/none.txt"})["error"]["code"] == "not_found"
        assert sandbox.helper("read", {"path": "/etc/passwd"})["error"]["code"] == "outside_workspace"
        assert sandbox.helper("read", {"path": "relative.txt"})["error"]["code"] == "invalid_argument"
        assert sandbox.helper("read", {"path": "/workspace/work"})["error"]["code"] == "not_a_file"


class TestWriteAndEdit:
    def test_write_a_new_file_then_overwrite_after_reading(self, sandbox):
        created = sandbox.helper("write", {"path": "/workspace/work/new/a.txt", "content": "あ\nい\n"})
        assert created["created"] is True
        again = sandbox.helper("write", {"path": "/workspace/work/new/a.txt", "content": "x"})
        assert again["error"]["code"] == "not_read"
        updated = sandbox.helper(
            "write", {"path": "/workspace/work/new/a.txt", "content": "う\n", "expected_sha256": created["sha256"]}
        )
        assert updated["created"] is False
        stale = sandbox.helper(
            "write", {"path": "/workspace/work/new/a.txt", "content": "え\n", "expected_sha256": created["sha256"]}
        )
        assert stale["error"]["code"] == "changed"

    def test_only_work_and_output_are_writable(self, sandbox):
        for path in ("/workspace/input/x.txt", "/workspace/master/x.txt", "/tmp/x.txt"):
            assert sandbox.helper("write", {"path": path, "content": ""})["error"]["code"] in (
                "read_only",
                "outside_workspace",
            ), path
        # symlinkを辿った先も確かめる
        sandbox.sh("ln -sf /workspace/input/メモ.txt /workspace/work/link.txt")
        assert sandbox.helper("write", {"path": "/workspace/work/link.txt", "content": ""})["error"]["code"] == (
            "read_only"
        )

    def test_edit_keeps_encoding_and_newlines(self, sandbox):
        sandbox.sh("cp /workspace/input/請求書一覧.csv /workspace/work/sjis.csv")
        read = sandbox.helper("read", {"path": "/workspace/work/sjis.csv"})
        edited = sandbox.helper(
            "edit",
            {
                "path": "/workspace/work/sjis.csv",
                "old_string": "2,髙橋商店,2500",
                "new_string": "2,髙橋商店,3000\n3,渡邉工業,100",
                "expected_sha256": read["sha256"],
            },
        )
        assert edited["replacements"] == 1
        assert "3,渡邉工業,100" in edited["snippet"]
        after = sandbox.helper("read", {"path": "/workspace/work/sjis.csv"})
        assert (after["encoding"], after["newline"], after["total_lines"]) == ("cp932", "crlf", 4)

    def test_edit_matching(self, sandbox):
        created = sandbox.helper("write", {"path": "/workspace/work/code.py", "content": "a = 1   \nb = 1\nc = 1\n"})
        sha = created["sha256"]
        multiple = sandbox.helper(
            "edit",
            {"path": "/workspace/work/code.py", "old_string": "= 1", "new_string": "= 2", "expected_sha256": sha},
        )
        assert multiple["error"]["code"] == "multiple_matches"
        assert multiple["error"]["detail"]["lines"] == [1, 2, 3]
        missing = sandbox.helper(
            "edit", {"path": "/workspace/work/code.py", "old_string": "d = 1", "new_string": "", "expected_sha256": sha}
        )
        assert missing["error"]["code"] == "no_match"
        # 行末の空白の違いは無視して一致させる
        flexible = sandbox.helper(
            "edit",
            {
                "path": "/workspace/work/code.py",
                "old_string": "a = 1\nb = 1",
                "new_string": "a = 3",
                "expected_sha256": sha,
            },
        )
        assert flexible["replacements"] == 1
        every = sandbox.helper(
            "edit",
            {
                "path": "/workspace/work/code.py",
                "old_string": "= ",
                "new_string": "== ",
                "replace_all": True,
                "expected_sha256": flexible["sha256"],
            },
        )
        assert every["replacements"] == 2
        assert sandbox.sh("cat /workspace/work/code.py") == "a == 3\nc == 1\n"


class TestSearch:
    def test_glob_newest_first(self, sandbox):
        sandbox.sh("mkdir -p g && touch -d '2026-01-01' g/old.csv && touch g/new.csv")
        result = sandbox.helper("glob", {"pattern": "**/*.csv", "path": "/workspace/work/g"})
        assert [e["path"] for e in result["entries"]] == ["/workspace/work/g/new.csv", "/workspace/work/g/old.csv"]
        assert sandbox.helper("glob", {"pattern": "../*"})["error"]["code"] == "invalid_argument"

    def test_grep_modes(self, sandbox):
        sandbox.sh("mkdir -p /workspace/work/s && printf '請求\\nなし\\n請求書\\n' > /workspace/work/s/a.txt")
        files = sandbox.helper("grep", {"pattern": "請求", "path": "/workspace/work/s"})
        assert files["files"] == [{"path": "/workspace/work/s/a.txt", "matches": 2}]
        lines = sandbox.helper("grep", {"pattern": "請求", "path": "/workspace/work/s", "output": "lines"})
        assert [(x["line"], x["text"]) for x in lines["lines"]] == [(1, "請求"), (3, "請求書")]
        count = sandbox.helper("grep", {"pattern": "請求", "path": "/workspace/work/s", "output": "count"})
        assert count["total_matches"] == 2
        bad = sandbox.helper("grep", {"pattern": "(", "path": "/workspace/work/s"})
        assert bad["error"]["code"] == "invalid_argument"

    def test_stat(self, sandbox):
        assert sandbox.helper("stat", {"path": "/workspace/work/none"}) == {
            "ok": True,
            "path": "/workspace/work/none",
            "exists": False,
        }
        stat = sandbox.helper("stat", {"path": "/workspace/input/メモ.txt"})
        assert (stat["type"], stat["kind"]) == ("file", "text")


# 追加するpackageの見本。中身は1つのモジュールだけのwheel（xlsxwriterの版を上書きする形で、先に読まれることも確かめる）
MAKE_WHEEL = r"""
python - <<'PY'
import base64, hashlib, zipfile
name, version = "xlsxwriter", "99.0"
files = {
    "xlsxwriter/__init__.py": b"__version__ = '99.0'\n",
    f"{name}-{version}.dist-info/METADATA": f"Metadata-Version: 2.1\nName: {name}\nVersion: {version}\n".encode(),
    f"{name}-{version}.dist-info/WHEEL": (
        b"Wheel-Version: 1.0\nGenerator: test\nRoot-Is-Purelib: true\nTag: py3-none-any\n"
    ),
}
record = []
for path, data in files.items():
    digest = base64.urlsafe_b64encode(hashlib.sha256(data).digest()).rstrip(b"=").decode()
    record.append(f"{path},sha256={digest},{len(data)}")
record.append(f"{name}-{version}.dist-info/RECORD,,")
files[f"{name}-{version}.dist-info/RECORD"] = ("\n".join(record) + "\n").encode()
with zipfile.ZipFile(f"/workspace/work/{name}-{version}-py3-none-any.whl", "w") as z:
    for path, data in files.items():
        z.writestr(path, data)
PY
"""


class TestEnv:
    def test_env_is_a_thin_layer_over_the_image_python(self, sandbox):
        assert sandbox.helper("init-env", {}) == {"ok": True, "created": False}
        size_kb = int(sandbox.sh("du -sk /workspace/env | cut -f1"))
        assert size_kb < 1024
        assert sandbox.sh("python -c 'import pandas; print(pandas.__file__)'").startswith("/usr/local/lib/python3.13/")

    def test_added_packages_go_to_the_env_and_are_read_first(self, sandbox):
        sandbox.sh(MAKE_WHEEL)
        before = sandbox.sh("python -c 'import xlsxwriter; print(xlsxwriter.__version__)'").strip()
        assert before != "99.0"
        # Runnerと同じく `python -m pip` で入れる（エージェントが呼ぶ pip は、入れる操作を断る）
        sandbox.sh("python -m pip install --no-index --no-deps /workspace/work/xlsxwriter-99.0-py3-none-any.whl")
        location = sandbox.sh("python -c 'import xlsxwriter; print(xlsxwriter.__version__, xlsxwriter.__file__)'")
        assert location.split() == ["99.0", "/workspace/env/lib/python3.13/site-packages/xlsxwriter/__init__.py"]
        # 補助コマンドは、追加したpackageの影響を受けない（イメージのPythonを -I で動かす）
        assert sandbox.helper("version", {})["ok"]
        sandbox.sh("python -m pip uninstall -y xlsxwriter")
        assert sandbox.sh("python -c 'import xlsxwriter; print(xlsxwriter.__version__)'").strip() == before
